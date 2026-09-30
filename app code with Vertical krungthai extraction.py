import io
import re
import datetime
from decimal import Decimal

import pandas as pd
import pdfplumber
import streamlit as st

st.set_page_config(page_title="Statement Extraction", page_icon="🏦", layout="wide")

NUM = r"\d{1,3}(?:,\d{3})*\.\d{2}"
ROW = re.compile(r"^(\d{2}/\d{2}/\d{2})\s+(.+?)\s+((?:" + NUM + r"\s+)+)(\d{1,4})\s*$")
TYPE = re.compile(r"^(.*?\([A-Z]+\))\s*(.*)$")
TIME = re.compile(r"^(\d{2}:\d{2})\s*(.*)$")


def to_dec(s):
    return Decimal(s.replace(",", ""))


def parse_krungthai(file_bytes):
    lines = []
    with pdfplumber.open(io.BytesIO(file_bytes)) as pdf:
        for page in pdf.pages:
            text = page.extract_text() or ""
            lines.extend(text.split("\n"))
    rows = []
    i = 0
    while i < len(lines):
        m = ROW.match(lines[i].strip())
        if not m:
            i += 1
            continue
        d, mid, nums, br = m.groups()
        values = [to_dec(x) for x in nums.split()]
        tm = TYPE.match(mid)
        typ, det = (tm.group(1), tm.group(2)) if tm else (mid, "")
        time_str, extra = "", ""
        if i + 1 < len(lines):
            t = TIME.match(lines[i + 1].strip())
            if t:
                time_str, extra = t.group(1), t.group(2)
                i += 1
        rows.append(
            {
                "date": d,
                "time": time_str,
                "type": typ.strip(),
                "details": (det + " " + extra).strip(),
                "amount": values[-2] if len(values) >= 2 else None,
                "balance": values[-1],
                "branch": int(br),
            }
        )
        i += 1
    prev = None
    out = []
    for r in rows:
        if prev is None:
            dep = "เข้า" in r["type"]
            prev = r["balance"] - r["amount"] if dep else r["balance"] + r["amount"]
        diff = r["balance"] - prev
        r["withdrawal"] = float(-diff) if diff < 0 else None
        r["deposit"] = float(diff) if diff > 0 else None
        prev = r["balance"]
        out.append(r)
    df = pd.DataFrame(out)
    if df.empty:
        return df

    def conv(s):
        dd, mm, yy = s.split("/")
        return datetime.date(int(yy) + 2500 - 543, int(mm), int(dd))

    df["date"] = df["date"].map(conv)
    df["balance"] = df["balance"].astype(float)
    df = df[["date", "time", "type", "details", "withdrawal", "deposit", "balance", "branch"]]
    df.columns = [
        "Date",
        "Time",
        "Transaction Type",
        "Details / Ref",
        "Withdrawal (THB)",
        "Deposit (THB)",
        "Balance (THB)",
        "Branch",
    ]
    return df


def build_excel(df):
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill

    wb = Workbook()
    ws = wb.active
    ws.title = "Transactions"
    ws.append(list(df.columns) + ["Balance Check"])
    for c in ws[1]:
        c.font = Font(name="Arial", bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor="1F4E78")
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    first = df.iloc[0]
    opening = first["Balance (THB)"] - (first["Deposit (THB)"] if pd.notna(first["Deposit (THB)"]) else 0) + (
        first["Withdrawal (THB)"] if pd.notna(first["Withdrawal (THB)"]) else 0
    )
    for idx, r in enumerate(df.itertuples(index=False), start=2):
        tm = None
        if r[1]:
            h, mi = map(int, r[1].split(":"))
            tm = datetime.time(h, mi)
        prev = "Summary!$B$3" if idx == 2 else f"G{idx - 1}"
        ws.append(
            [
                r[0],
                tm,
                r[2],
                r[3],
                None if pd.isna(r[4]) else r[4],
                None if pd.isna(r[5]) else r[5],
                r[6],
                r[7],
                f'=IF(ROUND({prev}-N(E{idx})+N(F{idx})-G{idx},2)=0,"OK","CHECK")',
            ]
        )
    last = len(df) + 1
    for row in ws.iter_rows(min_row=2, max_row=last):
        for c in row:
            c.font = Font(name="Arial", size=10)
        row[0].number_format = "dd/mm/yyyy"
        row[1].number_format = "hh:mm"
        for k in (4, 5, 6):
            row[k].number_format = "#,##0.00"
    for col, w in zip("ABCDEFGHI", [12, 8, 30, 42, 16, 16, 18, 8, 14]):
        ws.column_dimensions[col].width = w
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:I{last}"

    s = wb.create_sheet("Summary", 0)
    norm = Font(name="Arial")
    s["A1"] = "Statement Summary"
    s["A1"].font = Font(name="Arial", bold=True, size=14)
    items = [
        (3, "Opening balance (THB)", round(float(opening), 2)),
        (4, "Total withdrawals (THB)", f"=SUM(Transactions!E2:E{last})"),
        (5, "Total deposits (THB)", f"=SUM(Transactions!F2:F{last})"),
        (6, "Closing balance (calc)", "=B3-B4+B5"),
        (7, "Closing balance (statement)", f"=Transactions!G{last}"),
        (8, "Difference", "=ROUND(B6-B7,2)"),
        (9, "Withdrawal count", f"=COUNT(Transactions!E2:E{last})"),
        (10, "Deposit count", f"=COUNT(Transactions!F2:F{last})"),
        (11, "Rows failing balance check", f'=COUNTIF(Transactions!I2:I{last},"CHECK")'),
    ]
    for r, a, b in items:
        s.cell(r, 1, a).font = norm
        c = s.cell(r, 2, b)
        c.font = Font(name="Arial", color="0000FF") if r == 3 else norm
        c.number_format = "#,##0.00" if r <= 8 else "0"
    s.column_dimensions["A"].width = 34
    s.column_dimensions["B"].width = 20
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


PARSERS = {"Krungthai Statement Extraction": parse_krungthai}

with st.sidebar:
    st.title("🏦 Statement Tools")
    menu = st.radio("Menu", list(PARSERS.keys()), label_visibility="collapsed")

st.title(menu)

uploaded = st.file_uploader("Upload statement PDF", type=["pdf"])

if uploaded is not None:
    data = uploaded.getvalue()
    with st.spinner("Extracting..."):
        try:
            df = PARSERS[menu](data)
        except Exception as e:
            df = None
            st.error(f"Extraction failed: {e}")
    if df is not None:
        if df.empty:
            st.warning("No transactions found.")
        else:
            wd = df["Withdrawal (THB)"].sum()
            dp = df["Deposit (THB)"].sum()
            c1, c2, c3, c4 = st.columns(4)
            c1.metric("Transactions", f"{len(df):,}")
            c2.metric("Total withdrawals", f"{wd:,.2f}")
            c3.metric("Total deposits", f"{dp:,.2f}")
            c4.metric("Closing balance", f"{df['Balance (THB)'].iloc[-1]:,.2f}")
            st.dataframe(df, use_container_width=True, height=560)
            st.download_button(
                "Download Excel",
                data=build_excel(df),
                file_name=f"{uploaded.name.rsplit('.', 1)[0]}.xlsx",
                mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            )